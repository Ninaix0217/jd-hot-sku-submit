"""Read supplied workbooks; write a manifest without changing sources."""
import argparse
import hashlib
import json
from pathlib import Path
from urllib.parse import urlparse
import openpyxl


def identifier(value):
    if isinstance(value, bool):
        raise ValueError('Boolean is not a SKU')
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f'Non-integer SKU: {value!r}')
        value = int(value)
    result = str(value).strip()
    if not result.isdigit():
        raise ValueError(f'Invalid SKU: {result!r}')
    return result


def sheet(path, name):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return book, book[name] if name else book.worksheets[0]


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--inspect', type=Path)
    p.add_argument('--sku-file', type=Path)
    p.add_argument('--sku-sheet')
    p.add_argument('--sku-column', type=int)
    p.add_argument('--sku-start-row', type=int, default=2)
    p.add_argument('--link-file', type=Path)
    p.add_argument('--link-sheet')
    p.add_argument('--link-column', type=int)
    p.add_argument('--link-start-row', type=int, default=2)
    p.add_argument('--image', type=Path)
    p.add_argument('--out', type=Path)
    a = p.parse_args()
    if a.inspect:
        w = openpyxl.load_workbook(a.inspect, read_only=True, data_only=True)
        print(json.dumps([{'sheet': s.title, 'rows': s.max_row, 'columns': s.max_column,
                           'preview': list(s.iter_rows(min_row=1, max_row=3, values_only=True))}
                          for s in w], ensure_ascii=True, default=str))
        w.close()
        return
    if not all([a.sku_file, a.sku_column, a.link_file, a.link_column, a.image, a.out]):
        p.error('Extraction requires SKU/link files and explicit columns, image, and out')
    if min(a.sku_column, a.link_column, a.sku_start_row, a.link_start_row) < 1:
        p.error('Columns/rows are 1-based')
    if a.image.suffix.lower() not in ('.png', '.jpg', '.jpeg') or not a.image.is_file():
        p.error('Image must exist and use png/jpg/jpeg')
    if a.image.stat().st_size > 5 * 1024 * 1024:
        p.error('Image exceeds the observed 5 MB page limit')
    skus, links, seen_skus, seen_links = [], [], set(), set()
    w, s = sheet(a.sku_file, a.sku_sheet)
    for row, values in enumerate(s.values, 1):
        if row < a.sku_start_row:
            continue
        v = values[a.sku_column - 1] if len(values) >= a.sku_column else None
        if v is None or str(v).strip() == '':
            continue
        sku = identifier(v)
        if sku not in seen_skus:
            seen_skus.add(sku)
            skus.append({'row': row, 'sku': sku})
    w.close()
    w, s = sheet(a.link_file, a.link_sheet)
    for row, values in enumerate(s.values, 1):
        if row < a.link_start_row:
            continue
        v = values[a.link_column - 1] if len(values) >= a.link_column else None
        if v is None or str(v).strip() == '':
            continue
        url = str(v).strip()
        parsed = urlparse(url)
        if parsed.scheme not in ('http', 'https') or not parsed.netloc:
            raise ValueError(f'Invalid link at row {row}: {url!r}')
        if url not in seen_links:
            seen_links.add(url)
            links.append({'row': row, 'url': url})
    w.close()
    if not skus or not links:
        raise ValueError('SKU/link extraction is empty; verify sheets, columns, and start rows')
    pool_id = hashlib.sha256('\n'.join(x['url'] for x in links).encode()).hexdigest()
    out = {'poolId': pool_id, 'skus': skus, 'links': links, 'image': str(a.image.resolve()),
           'sku_file': str(a.sku_file.resolve()), 'link_file': str(a.link_file.resolve())}
    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'uniqueSkus': len(skus), 'links': len(links), 'poolId': pool_id}))


if __name__ == '__main__':
    main()
